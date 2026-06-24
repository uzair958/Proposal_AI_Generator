from pathlib import Path

from docx import Document
from langchain.tools import tool


@tool
def generate_proposal_docx(
    proposal_sections: dict,
):
    """
    Generate DOCX proposal document.
    """

    output_dir = Path(
        "generated_proposals"
    )

    output_dir.mkdir(
        exist_ok=True
    )

    file_path = (
        output_dir
        / "proposal.docx"
    )

    document = Document()

    for section, content in proposal_sections.items():

        document.add_heading(
            section.replace(
                "_",
                " "
            ).title(),
            level=1,
        )

        document.add_paragraph(
            str(content)
        )

    document.save(
        str(file_path)
    )

    return str(file_path)